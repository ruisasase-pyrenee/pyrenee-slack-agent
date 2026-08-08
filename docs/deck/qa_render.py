# -*- coding: utf-8 -*-
"""
Render each slide from the ACTUAL generated .pptx geometry (python-pptx) with PIL,
so layout defects in the real file are visible. Also reports out-of-bounds shapes
and text likely to overflow its box.
"""
import sys
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

PPTX = sys.argv[1] if len(sys.argv) > 1 else "qa.pptx"
SCALE = 105  # px per inch
JP = "/usr/share/fonts/opentype/ipafont-gothic/ipagp.ttf"

prs = Presentation(PPTX)
SW = prs.slide_width / 914400
SH = prs.slide_height / 914400
PW, PH = int(SW * SCALE), int(SH * SCALE)
print(f"slide: {SW:.2f} x {SH:.2f} in  ->  {PW}x{PH}px")

_fonts = {}
def font(sz):
    k = max(8, int(sz))
    if k not in _fonts:
        _fonts[k] = ImageFont.truetype(JP, k)
    return _fonts[k]

def rgb(c, default=(120, 130, 150)):
    try:
        if c and c.type is not None and c.rgb is not None:
            return tuple(bytes.fromhex(str(c.rgb)))
    except Exception:
        pass
    return default

def emu_px(v):
    return v / 914400 * SCALE

issues = []

# ── font check: the defect a pixel renderer can never see ────────────────
import zipfile as _zf
_z = _zf.ZipFile(PPTX)
_LATIN_ONLY = ("Cambria", "Calibri", "Arial", "Times New Roman", "Aptos")
for _n in _z.namelist():
    if _n.endswith(".xml") and "/slides/" in _n:
        _t = _z.read(_n).decode("utf-8")
        for _f in _LATIN_ONLY:
            _c = _t.count(f'<a:ea typeface="{_f}"')
            if _c:
                issues.append(f"{_n}: FONT — {_c} run(s) route Japanese to "
                              f"'{_f}', which has no CJK glyphs (mojibake)")

for idx, slide in enumerate(prs.slides, 1):
    # background
    bg = (10, 14, 31)
    img = Image.new("RGB", (PW, PH), bg)
    d = ImageDraw.Draw(img)
    boxes, shapes = [], []

    for sh in slide.shapes:
        if sh.left is None or sh.top is None:
            continue
        x, y = emu_px(sh.left), emu_px(sh.top)
        w, h = emu_px(sh.width or 0), emu_px(sh.height or 0)

        # out-of-bounds check on the real file
        if x < -1 or y < -1 or x + w > PW + 1 or y + h > PH + 1:
            issues.append(f"S{idx}: out of bounds ({sh.shape_type}) "
                          f"x={x/SCALE:.2f} y={y/SCALE:.2f} w={w/SCALE:.2f} h={h/SCALE:.2f}")

        st = str(sh.shape_type)

        if "PICTURE" in st:
            try:
                from io import BytesIO
                pic = Image.open(BytesIO(sh.image.blob)).convert("RGBA")
                pic = pic.resize((max(1, int(w)), max(1, int(h))))
                img.paste(pic, (int(x), int(y)), pic)   # honour alpha
            except Exception:
                d.rectangle([x, y, x + w, y + h], outline=(90, 100, 130), width=2)
            continue

        if "CHART" in st or "GRAPHIC" in st:
            d.rectangle([x, y, x + w, y + h], fill=(22, 29, 54), outline=(90, 200, 240), width=2)
            d.text((x + 10, y + 8), "[CHART]", font=font(15), fill=(95, 211, 243))
            continue

        # autoshape fill
        fill_col = None
        try:
            if sh.fill.type is not None and str(sh.fill.type) != "MSO_FILL_TYPE.BACKGROUND (5)":
                fill_col = rgb(sh.fill.fore_color, None)
        except Exception:
            pass
        if fill_col:
            shapes.append((x, y, w, h))
            if "ROUNDED" in st or "OVAL" in st or "ELLIPSE" in st:
                if "OVAL" in st or "ELLIPSE" in st:
                    d.ellipse([x, y, x + w, y + h], fill=fill_col)
                else:
                    d.rounded_rectangle([x, y, x + w, y + h], radius=8, fill=fill_col,
                                        outline=(42, 52, 83))
            else:
                d.rectangle([x, y, x + w, y + h], fill=fill_col)

        # text
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        txt = tf.text
        if not txt.strip():
            continue

        # first run's size/color/bold
        sz, col, bold, align = 14, (230, 235, 245), False, "l"
        for p in tf.paragraphs:
            if p.alignment is not None:
                a = str(p.alignment)
                align = "c" if "CENTER" in a else ("r" if "RIGHT" in a else "l")
            for r in p.runs:
                if r.font.size:
                    sz = r.font.size.pt
                if r.font.bold:
                    bold = True
                try:
                    if r.font.color and r.font.color.rgb:
                        col = tuple(bytes.fromhex(str(r.font.color.rgb)))
                except Exception:
                    pass
                break
            break

        f = font(sz * SCALE / 72.0)
        # wrap into the box
        maxw = max(10, w - 6)
        lines = []
        for para in txt.split("\n"):
            cur = ""
            for ch in para:
                if d.textlength(cur + ch, font=f) <= maxw:
                    cur += ch
                else:
                    lines.append(cur)
                    cur = ch
            lines.append(cur)

        lh = f.size * 1.22
        total = len(lines) * lh
        if total > h + 2:
            issues.append(f"S{idx}: TEXT OVERFLOW  \"{txt[:34]}...\"  "
                          f"needs {total/SCALE:.2f}in, box {h/SCALE:.2f}in")

        boxes.append((x, y, w, min(h, total), txt))

        ty = y + 2
        for ln in lines:
            tw = d.textlength(ln, font=f)
            tx = x + 3
            if align == "c":
                tx = x + (w - tw) / 2
            elif align == "r":
                tx = x + w - tw - 3
            d.text((tx, ty), ln, font=f, fill=col)
            ty += lh

    # ── overlap + edge-margin checks on real geometry ──────────────────
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            (ax, ay, aw, ah, at), (bx, by, bw, bh, bt) = boxes[i], boxes[j]
            ox = min(ax + aw, bx + bw) - max(ax, bx)
            oy = min(ay + ah, by + bh) - max(ay, by)
            if ox > 6 and oy > 6:                       # >0.06in on both axes
                issues.append(f"S{idx}: TEXT OVERLAP  \"{at[:20]}\" x \"{bt[:20]}\"  "
                              f"({ox/SCALE:.2f} x {oy/SCALE:.2f} in)")
    for (bx, by, bw, bh, bt) in boxes:
        for (sx, sy, sw, sh_) in shapes:
            ox = min(bx + bw, sx + sw) - max(bx, sx)
            oy = min(by + bh, sy + sh_) - max(by, sy)
            inside = bx >= sx - 2 and by >= sy - 2 and \
                     bx + bw <= sx + sw + 2 and by + bh <= sy + sh_ + 2
            if ox > 6 and oy > 6 and not inside:
                issues.append(f"S{idx}: TEXT ON SHAPE EDGE  \"{bt[:22]}\" "
                              f"straddles a card boundary ({ox/SCALE:.2f} x {oy/SCALE:.2f} in)")
    for (bx, by, bw, bh, bt) in boxes:
        if bx < 0.45 * SCALE or by < 0.25 * SCALE or \
           bx + bw > PW - 0.45 * SCALE or by + bh > PH - 0.2 * SCALE:
            issues.append(f"S{idx}: EDGE MARGIN  \"{bt[:24]}\" too close to slide edge")

    img.save(f"qa-{idx:02d}.jpg", quality=88)

print(f"\nrendered {len(prs.slides._sldIdLst)} slides")
print("\n=== ISSUES ===")
if issues:
    for i in issues:
        print(" -", i)
else:
    print("none")
